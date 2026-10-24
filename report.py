"""
report.py
Excel (.xlsx) rapor uretimi: ozet, primerler, problar, korunmusluk,
capraz reaksiyon matrisi ve PCR protokolu sayfalari.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")


def _write_table(ws, start_row, headers, rows, widths=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w

    return start_row + len(rows)


def _sheet_title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, size=9, color="666666")


def build_panel_workbook(results, cross_rows=None, settings=None):
    """
    results: her virüs için
      {
        "virus", "segment", "accession", "reference", "region_text",
        "strict": bool, "seq_len": int,
        "best": {forward, reverse, forward_tm, reverse_tm, forward_gc,
                 reverse_gc, amplicon_size, score, forward_start, reverse_end},
        "probe": {...} veya None,
        "cons_info": {...} veya None,
        "candidates": [...],
      }
    cross_rows: [{"primer_of","primer","target","match","risk"}]
    """
    wb = Workbook()
    meta = settings or {}
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    # ---------------- Özet ----------------
    ws = wb.active
    ws.title = "Özet"
    _sheet_title(ws, "Real-Time PCR Primer & Prob Tasarım Raporu", f"Oluşturulma: {stamp}")

    info_rows = [
        ("Hedef amplikon boyutu", f"{meta.get('amplicon', '-')} bp"),
        ("Primer Tm hedefi", "58-65 °C (qPCR 50 mM Na+, 1.5 mM Mg²⁺, 250 nM)"),
        ("Prob Tm hedefi", "66-73 °C (~69 °C)"),
        ("Korunmuşluk analizi", "Açık" if meta.get("use_cons") else "Kapalı"),
        ("Karşılaştırılan dizi sayısı", meta.get("n_seqs", "-")),
        ("Tasarım modu", "Katı kriterler" if all(r.get("strict") for r in results) else "Karma / gevşek"),
    ]
    r = 4
    for label, value in info_rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    headers = ["Virüs", "Segment", "Accession", "Bölge", "Amplikon (bp)",
               "F Tm", "R Tm", "Prob Tm", "Ort. korunmuşluk (%)", "Mod", "Skor"]
    rows = []
    for res in results:
        best = res.get("best") or {}
        probe = res.get("probe") or {}
        cons = res.get("cons_info") or {}
        rows.append([
            res.get("virus"), res.get("segment"), res.get("accession"),
            res.get("region_text"),
            best.get("amplicon_size"),
            best.get("forward_tm"), best.get("reverse_tm"), probe.get("tm"),
            round(cons.get("mean", 0) * 100, 1) if cons else "—",
            "Katı" if res.get("strict") else "Gevşek",
            best.get("score"),
        ])
    _write_table(ws, r, headers, rows, widths=[26, 28, 14, 18, 13, 8, 8, 9, 20, 9, 9])

    # ---------------- Primerler ----------------
    ws2 = wb.create_sheet("Primerler")
    _sheet_title(ws2, "Primer Dizileri")
    headers = ["Virüs", "Yön", "Dizi (5'→3')", "Uzunluk (nt)", "Tm (°C)", "GC (%)",
               "Konum (nt)", "Korunmuşluk (%)", "Skor"]
    rows = []
    for res in results:
        best = res.get("best") or {}
        cons = res.get("cons_info") or {}
        cons_vals = cons.get("conservation") or []
        if not best:
            continue
        f_cons = r_cons = "—"
        if cons_vals:
            off = max(0, (best.get("forward_start") or 1) - 1)
            f_cons = round(
                sum(cons_vals[off:off + best.get("forward_len", 20)]) /
                max(1, len(cons_vals[off:off + best.get("forward_len", 20)])), 3
            ) * 100
            r_end = best.get("reverse_end") or 0
            r_len = best.get("reverse_len", 20)
            seg = cons_vals[max(0, r_end - r_len):r_end]
            r_cons = round(sum(seg) / max(1, len(seg)), 3) * 100
        rows.append([res.get("virus"), "Forward", best.get("forward"),
                     best.get("forward_len"), best.get("forward_tm"),
                     best.get("forward_gc"), best.get("forward_start"), f_cons, best.get("score")])
        rows.append([res.get("virus"), "Reverse", best.get("reverse"),
                     best.get("reverse_len"), best.get("reverse_tm"),
                     best.get("reverse_gc"), best.get("reverse_end"), r_cons, best.get("score")])
    _write_table(ws2, 3, headers, rows, widths=[26, 11, 34, 12, 9, 8, 12, 16, 9])

    # ---------------- Problar ----------------
    ws3 = wb.create_sheet("Problar")
    _sheet_title(ws3, "TaqMan Prob Dizileri")
    headers = ["Virüs", "Prob (5'→3')", "Uzunluk (nt)", "Tm (°C)", "GC (%)",
               "Konum (nt)", "Korunmuşluk (%)", "Önerilen etiket"]
    rows = []
    for res in results:
        p = res.get("probe")
        if not p:
            rows.append([res.get("virus"), "Prob bulunamadı", "", "", "", "", "", ""])
            continue
        rows.append([
            res.get("virus"), p.get("sequence"), p.get("length"), p.get("tm"),
            p.get("gc"), p.get("position"), p.get("conservation"),
            "[FAM]-5' ... 3'-[BHQ1]",
        ])
    _write_table(ws3, 3, headers, rows, widths=[26, 34, 12, 9, 8, 12, 16, 26])

    # ---------------- Korunmuşluk ----------------
    ws4 = wb.create_sheet("Korunmuşluk")
    _sheet_title(ws4, "Korunmuşluk Analizi (çoklu hizalama)")
    r = 3
    for res in results:
        cons = res.get("cons_info")
        ws4.cell(row=r, column=1, value=res.get("virus")).font = Font(bold=True, size=12)
        r += 1
        if not cons:
            ws4.cell(row=r, column=1, value="Bu etken için korunmuşluk analizi yapılmadı.")
            r += 2
            continue
        ws4.cell(row=r, column=1, value="Kullanılan dizi sayısı")
        ws4.cell(row=r, column=2, value=cons.get("n_used"))
        r += 1
        ws4.cell(row=r, column=1, value="Ortalama korunmuşluk")
        ws4.cell(row=r, column=2, value=round(cons.get("mean", 0) * 100, 1))
        r += 2

        if cons.get("details"):
            headers = ["Accession", "Tanım", "Kimlik (%)", "Kapsam (%)"]
            rows = [[d["id"], (d.get("description") or "")[:60], d["identity"], d["coverage"]]
                    for d in cons["details"]]
            r = _write_table(ws4, r, headers, rows, widths=[18, 62, 12, 12]) + 2

        cons_vals = cons.get("conservation") or []
        if cons_vals:
            ws4.cell(row=r, column=1, value="Pozisyon bazlı korunmuşluk (1 = %100)").font = Font(bold=True)
            r += 1
            chunk = 12
            for i in range(0, len(cons_vals), chunk):
                ws4.cell(row=r, column=1, value=i + 1)
                for j, v in enumerate(cons_vals[i:i + chunk]):
                    cell = ws4.cell(row=r, column=2 + j, value=v)
                    if v is None:
                        continue
                    if v >= 0.95:
                        cell.fill = OK_FILL
                    elif v >= 0.85:
                        cell.fill = WARN_FILL
                    else:
                        cell.fill = BAD_FILL
                r += 1
        r += 2

    # ---------------- Çapraz Reaksiyon ----------------
    ws5 = wb.create_sheet("Capraz Reaksiyon")
    _sheet_title(ws5, "Multiplex Çapraz Reaksiyon Kontrolü",
                 "Primerlerin diğer panel etkenlerinin genomlarında aranması")
    if cross_rows is None:
        ws5["A3"] = "Çapraz reaksiyon taraması yapılmadı (tek etken modu)."
    elif not cross_rows:
        ws5["A3"] = "Seçilen panel etkenleri arasında primer eşleşmesi bulunmadı."
        ws5["A3"].font = Font(bold=True, color="006100")
    else:
        headers = ["Primer (ait olduğu etken)", "Primer dizisi", "Aranan genom", "Eşleşme", "Risk"]
        rows = [[c["primer_of"], c["primer"], c["target"], c["match"], c["risk"]]
                for c in cross_rows]
        _write_table(ws5, 3, headers, rows, widths=[26, 34, 28, 26, 12])
        ws5.freeze_panes = "A4"

    # ---------------- Protokol ----------------
    ws6 = wb.create_sheet("Protokol")
    _sheet_title(ws6, "Önerilen Real-Time PCR Protokolü (TaqMan, tek adım RT-PCR)")
    protocol = [
        ("Kalibrasyon", "Primer/prob konsantrasyonu: 200-400 nM primer, 100-250 nM prob"),
        ("Reaksiyon hacmi", "20 µL (10 µL 2x master mix + 1.5 µL primer mix + 0.5 µL prob + 5 µL RNA + su)"),
        ("RT adımı", "45 °C 10-15 dk"),
        ("İnaktivasyon", "95 °C 2-3 dk"),
        ("Döngü", "95 °C 5-15 sn → 60 °C 30-45 sn (45 döngü, FAM kanalı)"),
        ("Analiz", "Ct < 38 pozitif; NTC'de amplifikasyon olmamalı"),
        ("Doğrulama", "Primer-BLAST, dimer analizi (ör. OligoAnalyzer), LOD/LOQ çalışması"),
        ("Not", "Bu rapor in silico ön değerlendirmedir; klinik kullanım öncesi validasyon zorunludur."),
    ]
    r = 3
    for k, v in protocol:
        ws6.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = ws6.cell(row=r, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws6.column_dimensions["A"].width = 20
    ws6.column_dimensions["B"].width = 100

    # ---------------- Adaylar (varsa) ----------------
    all_candidates = []
    for res in results:
        for i, c in enumerate(res.get("candidates") or [], start=1):
            all_candidates.append([
                res.get("virus"), i, c.get("forward"), c.get("forward_tm"), c.get("forward_gc"),
                c.get("reverse"), c.get("reverse_tm"), c.get("reverse_gc"),
                c.get("amplicon_size"), c.get("forward_start"), c.get("reverse_end"),
                c.get("score"),
            ])
    if all_candidates:
        ws7 = wb.create_sheet("Adaylar")
        _sheet_title(ws7, "Tüm Aday Primer Setleri")
        headers = ["Virüs", "#", "Forward", "F Tm", "F GC", "Reverse", "R Tm", "R GC",
                   "Amplikon (bp)", "F başlangıç", "R bitiş", "Skor"]
        _write_table(ws7, 3, headers, all_candidates,
                     widths=[26, 5, 34, 8, 8, 34, 8, 8, 13, 13, 11, 9])
        ws7.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
